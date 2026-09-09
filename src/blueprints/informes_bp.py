"""Rutas de informes: histórico, KPIs e informe integral de desempeño."""
import csv
import io
from datetime import datetime
from pathlib import Path

from flask import Blueprint, render_template, request, Response, send_file, flash
from flask.typing import ResponseReturnValue
from historial import cargar_historial
from kpi_report import generar_informe_kpis
from dashboard_data import _filtrar, _rango_fechas, _en_rango
from drive_local import guardar_archivo_suelto_en_drive
from logging_config import obtener_logger
from informe_desempeno import (
    construir_informe, generar_excel, rango_predefinido, normalizar_nombre
)

log = obtener_logger("informes_bp")
informes_bp = Blueprint("informes", __name__)

BASE_DIR = Path(__file__).resolve().parent.parent.parent
SALIDAS_DIR = BASE_DIR / "salidas"


@informes_bp.route("/dashboard/exportar-excel")
def exportar_excel_del_dashboard() -> ResponseReturnValue:
    periodo = request.args.get("periodo", "todo")
    pais = request.args.get("pais", "").strip() or None
    bandeja = request.args.get("bandeja", "").strip() or None
    area = request.args.get("area", "").strip() or None
    fecha_desde = request.args.get("fecha_desde", "").strip() or None
    fecha_hasta = request.args.get("fecha_hasta", "").strip() or None
    todos = cargar_historial()
    base = _filtrar(todos, pais=pais, bandeja=bandeja, area=area)
    inicio, fin, _, _ = _rango_fechas(periodo, fecha_desde, fecha_hasta)
    registros = _en_rango(base, inicio, fin)

    partes_titulo = ["Informe de calidad"]
    if pais:
        partes_titulo.append(pais)
    if area:
        partes_titulo.append(area)
    if bandeja:
        partes_titulo.append(bandeja)
    titulo = " — ".join(partes_titulo)
    fecha_hoy = datetime.now().strftime("%Y-%m-%d")
    nombre_archivo = f"Dashboard {fecha_hoy}.xlsx"
    ruta_informe = SALIDAS_DIR / nombre_archivo
    generar_informe_kpis(registros, str(ruta_informe), titulo=titulo)

    drive_resultado = guardar_archivo_suelto_en_drive(str(ruta_informe))
    if drive_resultado.get("ok"):
        log.info("Excel del dashboard guardado en Drive: %s", drive_resultado["carpeta"])
    return send_file(ruta_informe, as_attachment=True, download_name=nombre_archivo)


@informes_bp.route("/informe")
def informe_kpis() -> ResponseReturnValue:
    alcance = request.args.get("alcance", "todo")
    solo_hoy = alcance == "hoy"
    registros = cargar_historial(solo_hoy=solo_hoy)
    fecha_hoy = datetime.now().strftime("%Y-%m-%d")
    if solo_hoy:
        titulo = f"Informe de evaluaciones del día ({fecha_hoy})"
        nombre_archivo = f"Informe de todas las conversaciones - Hoy {fecha_hoy}.xlsx"
    else:
        titulo = "Informe histórico de todas las evaluaciones"
        nombre_archivo = f"Informe de todas las conversaciones - Historico {fecha_hoy}.xlsx"
    ruta_informe = SALIDAS_DIR / nombre_archivo
    generar_informe_kpis(registros, str(ruta_informe), titulo=titulo)
    drive_resultado = guardar_archivo_suelto_en_drive(str(ruta_informe))
    return render_template(
        "informe_kpis.html", alcance=alcance, total=len(registros),
        nombre_archivo=nombre_archivo, drive_resultado=drive_resultado, activo="informes",
    )


@informes_bp.route("/historial.csv")
def exportar_historial_csv() -> ResponseReturnValue:
    registros = cargar_historial()
    salida = io.StringIO()
    escritor = csv.writer(salida)
    escritor.writerow(["Fecha", "Asesor", "Bandeja", "ID Caso", "País", "Evaluado por", "Nota Final (%)", "Ítem Crítico"])
    for r in sorted(registros, key=lambda x: x.get("timestamp", ""), reverse=True):
        escritor.writerow([
            r.get("fecha", ""), r.get("agente", ""), r.get("bandeja", ""), r.get("id_caso", ""),
            r.get("pais", ""), r.get("evaluado_por", ""),
            round(r.get("nota_final", 0) * 100, 1), r.get("critico_activado") or "",
        ])
    nombre_archivo = f"historial_evaluaciones_{datetime.now().strftime('%Y-%m-%d')}.csv"
    return Response(
        "\ufeff" + salida.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename={nombre_archivo}"},
    )


def _valor(nombre, default=""):
    if request.method == "POST":
        return request.form.get(nombre, default)
    return request.args.get(nombre, default)


def _lista_pegada(texto):
    nombres = []
    for linea in (texto or "").replace(";", "\n").splitlines():
        nombre = linea.strip().strip(",")
        if nombre:
            nombres.append(nombre)
    return nombres


def _nombres_desde_archivo(archivo):
    """Lee columna Asesor/Agente/Nombre de XLSX/XLSM/CSV. Fallback: primera columna."""
    if not archivo or not archivo.filename:
        return []
    nombre = archivo.filename.lower()

    if nombre.endswith(".csv"):
        contenido = archivo.read().decode("utf-8-sig", errors="replace")
        filas = list(csv.reader(io.StringIO(contenido)))
        if not filas:
            return []
        encabezados = [normalizar_nombre(x) for x in filas[0]]
        candidatos = {"asesor", "agente", "nombre", "advisor", "agent"}
        idx = next((i for i, h in enumerate(encabezados) if h in candidatos), 0)
        inicio = 1 if any(h in candidatos for h in encabezados) else 0
        return [str(f[idx]).strip() for f in filas[inicio:] if len(f) > idx and str(f[idx]).strip()]

    if nombre.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(archivo, read_only=True, data_only=True)
        ws = wb.active
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            return []
        encabezados = [normalizar_nombre(x) for x in filas[0]]
        candidatos = {"asesor", "agente", "nombre", "advisor", "agent"}
        idx = next((i for i, h in enumerate(encabezados) if h in candidatos), 0)
        inicio = 1 if any(h in candidatos for h in encabezados) else 0
        salida = []
        for fila in filas[inicio:]:
            if len(fila) > idx and fila[idx] is not None and str(fila[idx]).strip():
                salida.append(str(fila[idx]).strip())
        return salida

    raise ValueError("Formato no soportado. Usa .xlsx, .xlsm o .csv")


def _dedupe(nombres):
    salida, vistos = [], set()
    for n in nombres:
        n = str(n or "").strip()
        k = normalizar_nombre(n)
        if n and k and k not in vistos:
            vistos.add(k)
            salida.append(n)
    return salida


def _filtros_desempeno():
    periodo = _valor("periodo", "mes")
    minimo = _valor("minimo", "10")
    top_n = _valor("top_n", "10")
    modo = _valor("modo", "completo")

    try:
        minimo = max(1, min(int(minimo), 500))
    except (TypeError, ValueError):
        minimo = 10
    try:
        top_n = max(1, min(int(top_n), 500))
    except (TypeError, ValueError):
        top_n = 10

    fecha_desde = _valor("fecha_desde", "").strip() or None
    fecha_hasta = _valor("fecha_hasta", "").strip() or None
    if periodo != "personalizado":
        fecha_desde, fecha_hasta = rango_predefinido(periodo)

    asesores_texto = _valor("asesores_texto", "")
    seleccion_check = request.form.getlist("asesores_check") if request.method == "POST" else []
    seleccion = _lista_pegada(asesores_texto) + seleccion_check

    archivo = request.files.get("archivo_asesores") if request.method == "POST" else None
    if archivo and archivo.filename:
        try:
            seleccion.extend(_nombres_desde_archivo(archivo))
        except Exception as exc:
            flash(f"No se pudo leer el archivo de asesores: {exc}")

    seleccion = _dedupe(seleccion)

    return {
        "periodo": periodo,
        "minimo": minimo,
        "top_n": top_n,
        "modo": modo if modo in {"completo", "seleccion"} else "completo",
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "pais": _valor("pais", "").strip() or None,
        "bandeja": _valor("bandeja", "").strip() or None,
        "asesores_texto": "\n".join(seleccion),
        "asesores_seleccionados": seleccion,
    }


@informes_bp.route("/informe-desempeno", methods=["GET", "POST"])
def informe_desempeno() -> ResponseReturnValue:
    filtros = _filtros_desempeno()
    todos = cargar_historial()

    seleccion = filtros["asesores_seleccionados"] if filtros["modo"] == "seleccion" else None
    informe = construir_informe(
        todos,
        minimo_auditorias=filtros["minimo"],
        fecha_desde=filtros["fecha_desde"],
        fecha_hasta=filtros["fecha_hasta"],
        pais=filtros["pais"],
        bandeja=filtros["bandeja"],
        asesores_seleccionados=seleccion,
        top_n=filtros["top_n"],
    )

    paises = sorted({(r.get("pais") or "").strip() for r in todos if (r.get("pais") or "").strip()})
    bandejas = sorted({(r.get("bandeja") or "").strip() for r in todos if (r.get("bandeja") or "").strip()})
    asesores = sorted({(r.get("agente") or "").strip() for r in todos if (r.get("agente") or "").strip()})

    return render_template(
        "informe_desempeno.html",
        informe=informe, filtros=filtros, paises=paises, bandejas=bandejas,
        asesores=asesores, activo="informe_desempeno",
    )


@informes_bp.route("/informe-desempeno/excel", methods=["GET", "POST"])
def informe_desempeno_excel() -> ResponseReturnValue:
    filtros = _filtros_desempeno()
    seleccion = filtros["asesores_seleccionados"] if filtros["modo"] == "seleccion" else None
    informe = construir_informe(
        cargar_historial(),
        minimo_auditorias=filtros["minimo"],
        fecha_desde=filtros["fecha_desde"],
        fecha_hasta=filtros["fecha_hasta"],
        pais=filtros["pais"],
        bandeja=filtros["bandeja"],
        asesores_seleccionados=seleccion,
        top_n=filtros["top_n"],
    )
    fecha_hoy = datetime.now().strftime("%Y-%m-%d")
    sufijo = "_Grupo_Seleccionado" if filtros["modo"] == "seleccion" else ""
    nombre = f"Informe_Desempeno_QA{sufijo}_{fecha_hoy}.xlsx"
    ruta = SALIDAS_DIR / nombre
    generar_excel(informe, ruta)
    try:
        drive_resultado = guardar_archivo_suelto_en_drive(str(ruta))
        if drive_resultado.get("ok"):
            log.info("Informe de desempeño guardado en Drive: %s", drive_resultado.get("carpeta"))
    except Exception as exc:
        log.warning("No se pudo respaldar informe de desempeño en Drive: %s", exc)
    return send_file(ruta, as_attachment=True, download_name=nombre)
