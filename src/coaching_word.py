"""Genera el plan de coaching de un asesor como documento Word, en el mismo
estilo visual que el resto de informes (logo, naranja Dropi, bullets con
negrita), para poder imprimirlo o adjuntarlo a una conversación de desarrollo.
"""
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xF2, 0x65, 0x22)
GRIS = RGBColor(0x35, 0x30, 0x2B)
LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "dropi_logo.png"


def _shade_cell(cell: Any, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _titulo_seccion(doc: Any, texto: str) -> Any:
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = NARANJA
    p.space_after = Pt(4)
    return p


def _bullet(doc: Any, texto: str) -> None:
    doc.add_paragraph(texto, style="List Bullet")


def generar_word_coaching(asesor: str, datos: dict, plan: dict, ruta_salida: str) -> str:
    """Genera el plan de coaching como documento Word descargable, con el mismo estilo visual que los demás informes."""
    doc = Document()

    if LOGO_PATH.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(str(LOGO_PATH), width=Pt(110))

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(f"Plan de Coaching — {asesor}")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NARANJA

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub.add_run(
        f"Periodo analizado: {datos['primera_fecha']} a {datos['ultima_fecha']} "
        f"· {datos['total']} evaluaciones · Nota promedio: {round(datos['nota_promedio']*100,1)}%"
    )
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = GRIS
    doc.add_paragraph()

    seguimiento = plan.get("seguimiento_compromisos", "")
    if seguimiento and "no hay compromisos previos" not in seguimiento:
        _titulo_seccion(doc, "🔁 Seguimiento del compromiso anterior")
        p_seg = doc.add_paragraph(seguimiento)
        p_seg.runs[0].bold = True
        doc.add_paragraph()

    _titulo_seccion(doc, "Resumen general")
    doc.add_paragraph(plan.get("resumen_general", ""))
    p_tend = doc.add_paragraph()
    p_tend.add_run("Tendencia: ").bold = True
    p_tend.add_run(plan.get("tendencia", ""))
    doc.add_paragraph()

    _titulo_seccion(doc, "✅ Fortalezas consistentes")
    for f in plan.get("fortalezas_consistentes", []):
        _bullet(doc, f)
    doc.add_paragraph()

    _titulo_seccion(doc, "⚠ Áreas prioritarias")
    for area in plan.get("areas_prioritarias", []):
        p = doc.add_paragraph()
        run_nombre = p.add_run(f"{area.get('area','')} (impacto {area.get('impacto','')}): ")
        run_nombre.bold = True
        p.add_run(area.get("evidencia", ""))
    doc.add_paragraph()

    _titulo_seccion(doc, "📋 Plan de acción para la próxima conversación de desarrollo")
    for i, accion in enumerate(plan.get("plan_accion", []), start=1):
        doc.add_paragraph(f"{i}. {accion}")
    doc.add_paragraph()

    _titulo_seccion(doc, "Desempeño por categoría (periodo analizado)")
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = "Light Grid Accent 2"
    hdr = tabla.rows[0].cells
    hdr[0].text, hdr[1].text = "Categoría", "Eficiencia"
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
        _shade_cell(cell, "404040")
    for cat, val in datos["categorias_ordenadas"]:
        row = tabla.add_row().cells
        row[0].text = cat
        row[1].text = f"{round(val*100,1)}%"

    # --- Evolución real: mes actual vs. mes anterior — el ciclo natural de
    # coaching a este volumen (~40 evaluaciones/asesor/mes), y se mantiene
    # útil para siempre, sin importar cuántos meses de historial se acumulen ---
    if datos.get("mes_anterior"):
        doc.add_paragraph()
        _titulo_seccion(doc, "Evolución mes a mes")
        p_evidencia = doc.add_paragraph()
        p_evidencia.add_run(
            f"{datos['mes_anterior']} ({datos['cantidad_mes_anterior']} evaluaciones)   →   "
            f"{datos['mes_actual']} ({datos['cantidad_mes_actual']} evaluaciones)"
        ).italic = True

        tabla_evol = doc.add_table(rows=1, cols=4)
        tabla_evol.style = "Light Grid Accent 2"
        hdr2 = tabla_evol.rows[0].cells
        for i, texto in enumerate(["Categoría", f"Mes anterior ({datos['mes_anterior']})", f"Mes actual ({datos['mes_actual']})", "Cambio"]):
            hdr2[i].text = texto
            for p in hdr2[i].paragraphs:
                for r in p.runs:
                    r.bold = True
            _shade_cell(hdr2[i], "404040")

        cat_mes_anterior = datos.get("categorias_mes_anterior", {})
        cat_mes_actual = datos.get("categorias_mes_actual", {})
        for cat, _ in datos["categorias_ordenadas"]:
            pct_anterior = cat_mes_anterior.get(cat)
            pct_actual = cat_mes_actual.get(cat)
            if pct_anterior is None or pct_actual is None:
                continue
            fila = tabla_evol.add_row().cells
            fila[0].text = cat
            fila[1].text = f"{pct_anterior*100:.0f}%"
            fila[2].text = f"{pct_actual*100:.0f}%"
            diferencia = (pct_actual - pct_anterior) * 100
            simbolo = "▲" if diferencia > 0.5 else ("▼" if diferencia < -0.5 else "=")
            fila[3].text = f"{simbolo} {diferencia:+.0f} pts"
            color_cambio = RGBColor(0x1D, 0x7A, 0x4C) if diferencia > 0.5 else (RGBColor(0xC0, 0x00, 0x00) if diferencia < -0.5 else RGBColor(0x6B, 0x67, 0x59))
            fila[3].paragraphs[0].runs[0].font.color.rgb = color_cambio
    elif datos.get("mes_actual"):
        doc.add_paragraph()
        _titulo_seccion(doc, "Evolución mes a mes")
        p_sin_mes_anterior = doc.add_paragraph()
        p_sin_mes_anterior.add_run(
            f"Todas las evaluaciones de este asesor son de {datos['mes_actual']} — todavía no hay "
            "un mes anterior completo con el cual comparar la tendencia."
        ).italic = True

    # --- Firma: aquí es donde de verdad debe ir — el coordinador se compromete
    # a un plan de acción basado en la tendencia real del periodo, no en un
    # caso individual suelto ---
    doc.add_paragraph()
    _titulo_seccion(doc, "Compromiso de mejora")
    doc.add_paragraph(
        "El coordinador y el asesor revisaron juntos este plan de coaching, y se comprometen a dar "
        "seguimiento a las acciones acordadas antes de la próxima conversación de desarrollo."
    )
    doc.add_paragraph("_" * 70)
    doc.add_paragraph()

    p_firma_coord = doc.add_paragraph()
    p_firma_coord.add_run("Firma del coordinador: ").bold = True
    p_firma_coord.add_run("_" * 35)
    p_firma_coord.add_run("      Fecha: ").bold = True
    p_firma_coord.add_run("_" * 15)

    doc.add_paragraph()
    p_firma_asesor = doc.add_paragraph()
    p_firma_asesor.add_run("Firma del asesor: ").bold = True
    p_firma_asesor.add_run("_" * 35)
    p_firma_asesor.add_run("      Fecha: ").bold = True
    p_firma_asesor.add_run("_" * 15)

    doc.add_paragraph()
    nota_pie = doc.add_paragraph()
    run_pie = nota_pie.add_run("Generado por IA a partir del historial real de evaluaciones — revísalo antes de usarlo en la conversación con el asesor.")
    run_pie.italic = True
    run_pie.font.size = Pt(8)
    run_pie.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(ruta_salida)
    return ruta_salida
