"""Genera el Excel de la matriz de calidad rellenada con los puntajes de la evaluación,
replicando la estructura del archivo 'Matriz_Calidad_Dropi' original (misma disposición
de celdas, fórmula de nota final con lógica de críticos, formato visual similar).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

BASE_DIR = Path(__file__).resolve().parent.parent
MATRIZ_PATH = BASE_DIR / "config" / "matriz_calidad.json"
LOGO_PATH = BASE_DIR / "assets" / "dropi_logo.png"

NARANJA = "F26522"
GRIS_CLARO = "F2F2F2"
AMARILLO_CLARO = "FDEBD0"
VERDE_CLARO = "C6E0B4"
ROJO_CLARO = "F4CCCC"

THIN = Side(style="thin", color="BFBFBF")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cargar_matriz() -> dict:
    with open(MATRIZ_PATH, encoding="utf-8") as f:
        return json.load(f)


def generar_excel_matriz(evaluacion: dict, metadata: dict, ruta_salida: str) -> str:
    """
    evaluacion: dict con 'items' e 'items_criticos' (ver evaluator.py)
    metadata: dict con 'agente', 'fecha' (str), 'id_caso', 'bandeja', 'pais'
    ruta_salida: ruta del .xlsx a generar
    """
    matriz = _cargar_matriz()
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz de calidad"
    ws.sheet_view.showGridLines = False

    for col, width in zip("ABCDEF", [3, 42, 26, 10, 10, 12]):
        ws.column_dimensions[col].width = width

    # --- Logo Dropi, en el margen blanco arriba del banner ---
    if LOGO_PATH.exists():
        ws.row_dimensions[1].height = 32
        logo = XLImage(str(LOGO_PATH))
        logo.width = 95
        logo.height = 32
        ws.add_image(logo, "B1")

    # --- Encabezado ---
    ws.merge_cells("B2:F3")
    c = ws["B2"]
    c.value = "MATRIZ DE CALIDAD - DROPI"
    c.font = Font(size=16, bold=True, color="FFFFFF")
    c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws["B2:F3"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NARANJA)

    ws["B5"] = "Agente"
    ws["C5"] = metadata.get("agente", "")
    ws["B6"] = "Fecha"
    ws["C6"] = metadata.get("fecha", "")
    ws["B7"] = "ID / Bandeja"
    ws["C7"] = f"{metadata.get('id_caso', '')} - {metadata.get('bandeja', '')}"
    ws["B8"] = "País"
    ws["C8"] = metadata.get("pais", "")
    for r in range(5, 9):
        ws[f"B{r}"].font = Font(bold=True)

    ws["E5"] = "NOTA"
    ws["E5"].font = Font(bold=True, size=11)
    ws["E5"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E6"] = "FINAL"
    ws["E6"].font = Font(bold=True, size=11)
    ws["E6"].alignment = Alignment(horizontal="center", vertical="center")

    fila_items_inicio = 12
    fila = fila_items_inicio
    filas_items = []
    filas_criticos = []

    ws[f"B{fila}"] = "CATEGORÍA / ÍTEM"
    ws[f"E{fila}"] = "PESO MÁX"
    ws[f"F{fila}"] = "PUNTAJE OTORGADO"
    for col in "BEF":
        cell = ws[f"{col}{fila}"]
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="404040")
        cell.alignment = Alignment(horizontal="center" if col != "B" else "left")
    fila += 2

    items_eval = evaluacion.get("items", {})

    for cat in matriz["categorias"]:
        ws[f"B{fila}"] = cat["nombre"]
        ws[f"E{fila}"] = f"{cat['peso_categoria']*100:.0f}%"
        for col in "BCDEF":
            ws[f"{col}{fila}"].fill = PatternFill("solid", fgColor=GRIS_CLARO)
            ws[f"{col}{fila}"].font = Font(bold=True)
        fila += 1
        for it in cat["items"]:
            ws[f"B{fila}"] = it["nombre"]
            ws[f"B{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws[f"E{fila}"] = it["peso"]
            ws[f"E{fila}"].number_format = "0%"
            puntaje = items_eval.get(it["id"], {}).get("puntaje", 0)
            ws[f"F{fila}"] = round(float(puntaje), 4)
            ws[f"F{fila}"].number_format = "0.0%"
            justificacion = items_eval.get(it["id"], {}).get("justificacion", "")
            ws[f"C{fila}"] = justificacion
            ws[f"C{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[fila].height = 30
            for col in "BCEF":
                ws[f"{col}{fila}"].border = BORDE
            filas_items.append(fila)
            fila += 1
        fila += 1

    fila += 1
    ws[f"B{fila}"] = 'ÍTEMS CRÍTICOS (con al menos un "Sí", la nota final es 0%)'
    ws[f"F{fila}"] = "¿Ocurrió?"
    for col in "BEF":
        ws[f"{col}{fila}"].font = Font(bold=True, color="FFFFFF")
        ws[f"{col}{fila}"].fill = PatternFill("solid", fgColor="C00000")
    fila += 1

    criticos_eval = evaluacion.get("items_criticos", {})
    for c in matriz["items_criticos"]:
        ws[f"B{fila}"] = c["nombre"]
        ocurrio = criticos_eval.get(c["id"], {}).get("ocurrio", "No")
        ws[f"F{fila}"] = ocurrio
        ws[f"F{fila}"].alignment = Alignment(horizontal="center")
        ws[f"F{fila}"].fill = PatternFill(
            "solid", fgColor=(ROJO_CLARO if str(ocurrio).lower() == "si" else VERDE_CLARO)
        )
        justificacion = criticos_eval.get(c["id"], {}).get("justificacion", "")
        ws[f"C{fila}"] = justificacion
        ws[f"C{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
        for col in "BCF":
            ws[f"{col}{fila}"].border = BORDE
        filas_criticos.append(fila)
        fila += 1

    suma_items = "+".join(f"IFERROR(VALUE(F{f}),0)" for f in filas_items)
    condiciones_criticas = "".join(f'IF(F{f}="Si",0,' for f in filas_criticos)
    cierre_parentesis = ")" * len(filas_criticos)
    formula = f"={condiciones_criticas}({suma_items}){cierre_parentesis}"
    ws["F5"] = formula
    ws.merge_cells("F5:F6")
    ws["F5"].number_format = "0.0%"
    ws["F5"].font = Font(bold=True, size=20, color=NARANJA)
    ws["F5"].alignment = Alignment(horizontal="center", vertical="center")

    fila += 1
    ws[f"B{fila}"] = "Observaciones generales:"
    ws[f"B{fila}"].font = Font(bold=True)
    fila += 1
    ws.merge_cells(f"B{fila}:F{fila+3}")
    ws[f"B{fila}"] = evaluacion.get("resumen_caso", "")
    ws[f"B{fila}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[fila].height = 60

    wb.save(ruta_salida)
    return ruta_salida


if __name__ == "__main__":
    import sys
    with open(sys.argv[1], encoding="utf-8") as f:
        evaluacion = json.load(f)
    metadata = {
        "agente": "Stephany García",
        "fecha": "2026-07-27",
        "id_caso": "215475175505780",
        "bandeja": "Garantías",
        "pais": "Colombia",
    }
    ruta = generar_excel_matriz(evaluacion, metadata, "salidas/matriz_prueba.xlsx")
    print("Generado:", ruta)