"""Informe integral de desempeño QA por asesor y región.

Usa exclusivamente el historial confirmado de Dropi QA; no llama a ningún
proveedor de IA. Permite identificar:
- ranking global de asesores;
- peor calificado global (con mínimo de muestra);
- peor asesor por región/país;
- regiones con menor desempeño;
- casos críticos, reincidencias y tendencia;
- principal categoría y subítem de pérdida por asesor;
- índice de riesgo QA complementario.

También puede ejecutarse desde Terminal para generar el informe actual:
    python3 src/informe_desempeno.py --excel salidas/Informe_Desempeno_QA.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import mean

from historial import cargar_historial

BASE_DIR = Path(__file__).resolve().parent.parent
MATRIZ_PATH = BASE_DIR / "config" / "matriz_calidad.json"


def _numero(valor, default=0.0):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return default


def _nota_01(valor):
    valor = _numero(valor, 0.0)
    # Compatibilidad por si alguna fuente histórica guardó 0-100.
    if valor > 1.0001:
        valor /= 100.0
    return max(0.0, min(valor, 1.0))


def _fecha(registro):
    raw = registro.get("fecha") or ""
    try:
        return datetime.strptime(raw[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def _cargar_mapa_matriz():
    try:
        matriz = json.loads(MATRIZ_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}, {}
    items = {}
    categorias = {}
    for categoria in matriz.get("categorias", []):
        nombre_cat = categoria.get("nombre") or "Sin categoría"
        peso_cat = sum(_numero(i.get("peso"), 0) for i in categoria.get("items", []))
        categorias[nombre_cat] = peso_cat
        for item in categoria.get("items", []):
            iid = item.get("id")
            if not iid:
                continue
            items[iid] = {
                "id": iid,
                "nombre": item.get("nombre") or iid,
                "peso": _numero(item.get("peso"), 0),
                "categoria": nombre_cat,
            }
    return items, categorias


def filtrar_registros(registros, fecha_desde=None, fecha_hasta=None, pais=None, bandeja=None):
    desde = None
    hasta = None
    if fecha_desde:
        try:
            desde = datetime.strptime(fecha_desde, "%Y-%m-%d").date()
        except ValueError:
            pass
    if fecha_hasta:
        try:
            hasta = datetime.strptime(fecha_hasta, "%Y-%m-%d").date()
        except ValueError:
            pass

    salida = []
    for r in registros:
        if pais and (r.get("pais") or "").strip() != pais:
            continue
        if bandeja and (r.get("bandeja") or "").strip() != bandeja:
            continue
        d = _fecha(r)
        if desde and (not d or d < desde):
            continue
        if hasta and (not d or d > hasta):
            continue
        salida.append(r)
    return salida


def rango_predefinido(periodo):
    hoy = date.today()
    if periodo == "mes":
        return hoy.replace(day=1).isoformat(), hoy.isoformat()
    if periodo == "30d":
        return (hoy - timedelta(days=29)).isoformat(), hoy.isoformat()
    if periodo == "90d":
        return (hoy - timedelta(days=89)).isoformat(), hoy.isoformat()
    return None, None


def _detalle_item_score(raw):
    if isinstance(raw, dict):
        if raw.get("aplica") is False:
            return None
        raw = raw.get("puntaje")
    if raw is None:
        return None
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


def _analisis_items(registros, mapa_items, umbral=0.80):
    stats = defaultdict(lambda: {
        "evaluado": 0, "fallas": 0, "obtenido": 0.0,
        "maximo": 0.0, "perdida": 0.0,
    })
    for r in registros:
        for iid, raw in (r.get("items_detalle") or {}).items():
            info = mapa_items.get(iid)
            if not info or info["peso"] <= 0:
                continue
            score = _detalle_item_score(raw)
            if score is None:
                continue
            score = max(0.0, min(score, info["peso"]))
            s = stats[iid]
            s["evaluado"] += 1
            s["obtenido"] += score
            s["maximo"] += info["peso"]
            perdida = max(info["peso"] - score, 0.0)
            s["perdida"] += perdida
            if score / info["peso"] < umbral:
                s["fallas"] += 1

    filas = []
    for iid, s in stats.items():
        info = mapa_items[iid]
        cumplimiento = s["obtenido"] / s["maximo"] if s["maximo"] else 0
        filas.append({
            **info,
            "veces_evaluado": s["evaluado"],
            "fallas": s["fallas"],
            "tasa_falla": s["fallas"] / s["evaluado"] if s["evaluado"] else 0,
            "cumplimiento": cumplimiento,
            "puntos_perdidos": s["perdida"],
        })
    filas.sort(key=lambda x: (-x["puntos_perdidos"], -x["fallas"], x["nombre"]))
    return filas


def _categoria_peor(registros, categorias_matriz):
    # El historial ya guarda subtotales ponderados por categoría.
    # Comparamos cumplimiento = puntos obtenidos / peso de la categoría.
    datos = defaultdict(list)
    for r in registros:
        for nombre, valor in (r.get("categorias") or {}).items():
            maximo = _numero(categorias_matriz.get(nombre), 0)
            if maximo <= 0:
                continue
            score = _numero(valor, None)
            if score is None:
                continue
            datos[nombre].append(max(0.0, min(score / maximo, 1.0)))
    if not datos:
        return None
    nombre, valores = min(datos.items(), key=lambda kv: mean(kv[1]))
    return {"nombre": nombre, "cumplimiento": mean(valores), "evaluaciones": len(valores)}


def _tendencia(registros):
    ordenados = sorted(
        [r for r in registros if _fecha(r)],
        key=lambda r: (_fecha(r), r.get("timestamp") or "")
    )
    if len(ordenados) < 4:
        return {"delta": None, "direccion": "Sin muestra suficiente"}
    # Compara dos mitades cronológicas de la muestra del período.
    corte = len(ordenados) // 2
    anterior = [_nota_01(r.get("nota_final")) for r in ordenados[:corte]]
    reciente = [_nota_01(r.get("nota_final")) for r in ordenados[corte:]]
    delta = mean(reciente) - mean(anterior)
    if delta > 0.015:
        direccion = "Mejora"
    elif delta < -0.015:
        direccion = "Deterioro"
    else:
        direccion = "Estable"
    return {"delta": delta, "direccion": direccion}


def _riesgo_qa(promedio, tasa_criticos, criterios_reincidentes, delta_tendencia):
    """Índice complementario 0-100. No reemplaza el ranking por nota."""
    calidad = (1 - promedio) * 60
    criticos = min(max(tasa_criticos, 0) / 0.25, 1) * 20
    reincidencia = min(criterios_reincidentes / 3, 1) * 15
    deterioro = 0
    if delta_tendencia is not None and delta_tendencia < 0:
        deterioro = min(abs(delta_tendencia) / 0.15, 1) * 5
    return round(min(calidad + criticos + reincidencia + deterioro, 100), 1)


def _nivel_riesgo(indice):
    if indice >= 65:
        return "ALTO"
    if indice >= 40:
        return "MEDIO"
    return "BAJO"



def normalizar_nombre(nombre):
    """Normaliza nombres para cruces tolerantes a mayúsculas, tildes y espacios."""
    texto = unicodedata.normalize("NFKD", str(nombre or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.casefold().strip()
    texto = re.sub(r"\s+", " ", texto)
    return texto


def estado_muestra(auditorias):
    """Etiqueta descriptiva; no excluye asesores en modo de grupo seleccionado."""
    if auditorias <= 0:
        return {"codigo": "sin_datos", "texto": "Sin auditorías", "nivel": "gris"}
    if auditorias <= 4:
        return {"codigo": "muy_baja", "texto": "Muestra muy baja", "nivel": "rojo"}
    if auditorias <= 9:
        return {"codigo": "baja", "texto": "Muestra baja", "nivel": "ambar"}
    if auditorias <= 19:
        return {"codigo": "moderada", "texto": "Muestra moderada", "nivel": "ambar"}
    return {"codigo": "solida", "texto": "Muestra sólida", "nivel": "verde"}


def _deduplicar_nombres(nombres):
    salida = []
    vistos = set()
    for nombre in nombres or []:
        limpio = str(nombre or "").strip()
        clave = normalizar_nombre(limpio)
        if not limpio or not clave or clave in vistos:
            continue
        vistos.add(clave)
        salida.append(limpio)
    return salida

def construir_informe(registros=None, minimo_auditorias=10,
                      fecha_desde=None, fecha_hasta=None,
                      pais=None, bandeja=None,
                      asesores_seleccionados=None, top_n=None):
    """Construye el informe.

    Modo población completa:
      - aplica ``minimo_auditorias`` al ranking oficial.

    Modo grupo seleccionado:
      - analiza exclusivamente los nombres recibidos;
      - NO excluye por tamaño de muestra;
      - asesores con 1, 2, etc. auditorías entran al ranking;
      - asesores sin auditorías se conservan en ``sin_datos``;
      - ``estado_muestra`` advierte la robustez estadística;
      - ``top_n`` limita la tabla principal a los N peores, sin perder
        la población completa en ``ranking_completo``.
    """
    registros = cargar_historial() if registros is None else list(registros)
    registros = filtrar_registros(
        registros, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta,
        pais=pais, bandeja=bandeja
    )
    mapa_items, categorias_matriz = _cargar_mapa_matriz()

    seleccion = _deduplicar_nombres(asesores_seleccionados)
    modo_seleccion = bool(seleccion)

    # Índice de nombres existentes en el historial filtrado.
    nombres_historial = {}
    for r in registros:
        asesor = (r.get("agente") or "").strip()
        if asesor:
            nombres_historial.setdefault(normalizar_nombre(asesor), asesor)

    # En modo seleccionado, resolvemos nombres tolerando tildes/case/espacios.
    resolucion = []
    seleccion_resuelta = set()
    sin_datos = []
    if modo_seleccion:
        for solicitado in seleccion:
            clave = normalizar_nombre(solicitado)
            real = nombres_historial.get(clave)
            if real:
                resolucion.append({
                    "solicitado": solicitado, "encontrado": True,
                    "asesor": real, "coincidencia": "exacta_normalizada",
                })
                seleccion_resuelta.add(normalizar_nombre(real))
            else:
                resolucion.append({
                    "solicitado": solicitado, "encontrado": False,
                    "asesor": solicitado, "coincidencia": None,
                })
                sin_datos.append({
                    "asesor": solicitado,
                    "region": "Sin información",
                    "auditorias": 0,
                    "promedio": None,
                    "estado_muestra": estado_muestra(0),
                })

        registros_analisis = [
            r for r in registros
            if normalizar_nombre((r.get("agente") or "").strip()) in seleccion_resuelta
        ]
    else:
        registros_analisis = registros

    por_asesor = defaultdict(list)
    for r in registros_analisis:
        asesor = (r.get("agente") or "").strip()
        if asesor:
            por_asesor[asesor].append(r)

    elegibles = []
    insuficientes = []
    for asesor, regs in por_asesor.items():
        notas = [_nota_01(r.get("nota_final")) for r in regs]
        promedio = mean(notas) if notas else 0
        criticos = sum(1 for r in regs if r.get("critico_activado"))
        regiones = [str(r.get("pais") or "").strip() for r in regs if str(r.get("pais") or "").strip()]
        region = max(set(regiones), key=regiones.count) if regiones else "Sin región"
        bandejas = [str(r.get("bandeja") or "").strip() for r in regs if str(r.get("bandeja") or "").strip()]
        bandeja_principal = max(set(bandejas), key=bandejas.count) if bandejas else "Sin bandeja"

        items = _analisis_items(regs, mapa_items)
        recurrentes = [x for x in items if x["fallas"] >= 2]
        peor_item = items[0] if items else None
        peor_cat = _categoria_peor(regs, categorias_matriz)
        tendencia = _tendencia(regs)
        riesgo = _riesgo_qa(
            promedio,
            criticos / len(regs) if regs else 0,
            len(recurrentes),
            tendencia["delta"]
        )
        fila = {
            "asesor": asesor,
            "region": region,
            "bandeja_principal": bandeja_principal,
            "auditorias": len(regs),
            "promedio": promedio,
            "criticos": criticos,
            "tasa_criticos": criticos / len(regs) if regs else 0,
            "criterios_reincidentes": len(recurrentes),
            "fallas_reincidentes": sum(x["fallas"] for x in recurrentes),
            "tendencia_delta": tendencia["delta"],
            "tendencia": tendencia["direccion"],
            "peor_categoria": peor_cat,
            "peor_item": peor_item,
            "riesgo_qa": riesgo,
            "nivel_riesgo": _nivel_riesgo(riesgo),
            "estado_muestra": estado_muestra(len(regs)),
        }
        if modo_seleccion or len(regs) >= minimo_auditorias:
            elegibles.append(fila)
        else:
            insuficientes.append(fila)

    # Transparente: menor promedio QA = peor calificado.
    elegibles.sort(key=lambda x: (x["promedio"], -x["auditorias"], x["asesor"].lower()))
    for i, fila in enumerate(elegibles, 1):
        fila["posicion"] = i
    insuficientes.sort(key=lambda x: (x["promedio"], x["asesor"].lower()))

    try:
        top_n_int = int(top_n) if top_n not in (None, "", 0, "0") else None
        if top_n_int is not None:
            top_n_int = max(1, min(top_n_int, 500))
    except (TypeError, ValueError):
        top_n_int = None

    ranking_visible = elegibles[:top_n_int] if top_n_int else elegibles

    # Región calculada solo con la población analizada.
    por_region_regs = defaultdict(list)
    for r in registros_analisis:
        region = (r.get("pais") or "").strip() or "Sin región"
        por_region_regs[region].append(r)

    regiones = []
    peor_por_region = []
    for region, regs in por_region_regs.items():
        notas = [_nota_01(r.get("nota_final")) for r in regs]
        asesores_region = sorted({(r.get("agente") or "").strip() for r in regs if (r.get("agente") or "").strip()})
        criticos = sum(1 for r in regs if r.get("critico_activado"))
        regiones.append({
            "region": region,
            "auditorias": len(regs),
            "asesores": len(asesores_region),
            "promedio": mean(notas) if notas else 0,
            "criticos": criticos,
            "tasa_criticos": criticos / len(regs) if regs else 0,
        })
        candidatos = [x for x in elegibles if x["region"] == region]
        if candidatos:
            peor_por_region.append({
                "region": region,
                "asesor": candidatos[0]["asesor"],
                "promedio": candidatos[0]["promedio"],
                "auditorias": candidatos[0]["auditorias"],
                "riesgo_qa": candidatos[0]["riesgo_qa"],
            })

    regiones.sort(key=lambda x: (x["promedio"], -x["auditorias"], x["region"]))
    peor_por_region.sort(key=lambda x: (x["promedio"], x["region"]))

    perdida_global = _analisis_items(registros_analisis, mapa_items)[:10]

    return {
        "generado_en": datetime.now().isoformat(timespec="seconds"),
        "modo": "seleccion" if modo_seleccion else "completo",
        "filtros": {
            "fecha_desde": fecha_desde,
            "fecha_hasta": fecha_hasta,
            "pais": pais,
            "bandeja": bandeja,
            "minimo_auditorias": int(minimo_auditorias),
            "top_n": top_n_int,
        },
        "total_auditorias": len(registros_analisis),
        "total_asesores": len(por_asesor),
        "asesores_clasificados": len(elegibles),
        "ranking": ranking_visible,
        "ranking_completo": elegibles,
        "muestra_insuficiente": insuficientes,
        "peor_global": elegibles[0] if elegibles else None,
        "regiones": regiones,
        "peor_region": regiones[0] if regiones else None,
        "peor_por_region": peor_por_region,
        "perdida_global": perdida_global,
        "poblacion_solicitada": seleccion,
        "resolucion_seleccion": resolucion,
        "sin_datos": sin_datos,
        "total_solicitados": len(seleccion) if modo_seleccion else None,
        "total_encontrados": len(resolucion) - len(sin_datos) if modo_seleccion else None,
    }

def generar_excel(informe, ruta):
    """Genera un informe ejecutivo QA listo para presentar.

    El libro separa la lectura gerencial (Resumen ejecutivo) del detalle
    analítico. Usa el logo local de Dropi cuando está disponible.
    """
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.drawing.image import Image
    from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)

    # Paleta visual Dropi QA
    NARANJA = "F26522"
    NARANJA_OSCURO = "D9531A"
    NEGRO = "17140F"
    GRIS_OSCURO = "3A362E"
    GRIS = "6B6759"
    GRIS_CLARO = "F1F2EE"
    LINEA = "DAD6CC"
    BLANCO = "FFFFFF"
    VERDE = "24443E"
    VERDE_CLARO = "E4EBE9"
    AMBAR = "9C6B12"
    AMBAR_CLARO = "F3EAD6"
    ROJO = "B8342A"
    ROJO_CLARO = "F5E4E1"

    thin = Side(style="thin", color=LINEA)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def estilizar_titulo(ws, rango, texto):
        ws.merge_cells(rango)
        c = ws[rango.split(":")[0]]
        c.value = texto
        c.fill = PatternFill("solid", fgColor=NEGRO)
        c.font = Font(color=BLANCO, bold=True, size=18)
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[c.row].height = 30

    def encabezado_tabla(ws, fila, inicio, fin):
        for col in range(inicio, fin + 1):
            c = ws.cell(fila, col)
            c.fill = PatternFill("solid", fgColor=GRIS_OSCURO)
            c.font = Font(color=BLANCO, bold=True, size=10)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = border
        ws.row_dimensions[fila].height = 26

    def dar_formato_tabla(ws, fila_inicio, fila_fin, col_inicio, col_fin):
        for row in ws.iter_rows(
            min_row=fila_inicio, max_row=fila_fin,
            min_col=col_inicio, max_col=col_fin
        ):
            for c in row:
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in range(fila_inicio, fila_fin + 1):
            if (r - fila_inicio) % 2 == 1:
                for c in ws[r][col_inicio-1:col_fin]:
                    c.fill = PatternFill("solid", fgColor="FAFAF8")

    def kpi(ws, rango, titulo, valor, tipo="neutral"):
        ws.merge_cells(rango)
        c = ws[rango.split(":")[0]]
        colores = {
            "naranja": (NARANJA, BLANCO),
            "verde": (VERDE_CLARO, VERDE),
            "ambar": (AMBAR_CLARO, AMBAR),
            "rojo": (ROJO_CLARO, ROJO),
            "neutral": (GRIS_CLARO, NEGRO),
        }
        fill, font_color = colores.get(tipo, colores["neutral"])
        c.value = f"{titulo}\n{valor}"
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = Font(color=font_color, bold=True, size=13)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(
            left=Side(style="medium", color=fill if tipo != "neutral" else LINEA),
            right=thin, top=thin, bottom=thin
        )
        for row in ws[rango]:
            for cell in row:
                cell.border = border

    def formato_muestra(nivel):
        return {
            "verde": VERDE_CLARO,
            "ambar": AMBAR_CLARO,
            "rojo": ROJO_CLARO,
            "gris": GRIS_CLARO,
        }.get(nivel, GRIS_CLARO)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen ejecutivo"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A18"

    # Anchos del dashboard ejecutivo
    widths = {
        "A": 6, "B": 25, "C": 17, "D": 16, "E": 17,
        "F": 17, "G": 18, "H": 24, "I": 4,
        "J": 18, "K": 18, "L": 18, "M": 18, "N": 18,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # Logo
    logo = BASE_DIR / "static" / "dropi_logo.png"
    if logo.exists():
        try:
            img = Image(str(logo))
            img.width = 177
            img.height = 60
            ws.add_image(img, "A1")
        except Exception:
            pass

    ws.merge_cells("D1:N2")
    ws["D1"] = "INFORME INTEGRAL DE DESEMPEÑO QA"
    ws["D1"].font = Font(size=22, bold=True, color=NEGRO)
    ws["D1"].alignment = Alignment(vertical="center")
    ws.merge_cells("D3:N3")
    modo_txt = "Grupo seleccionado" if informe.get("modo") == "seleccion" else "Población completa"
    periodo = informe.get("filtros", {})
    desde = periodo.get("fecha_desde") or "Inicio"
    hasta = periodo.get("fecha_hasta") or "Actualidad"
    ws["D3"] = f"{modo_txt} · Período: {desde} a {hasta} · Generado: {informe.get('generado_en','')[:16].replace('T',' ')}"
    ws["D3"].font = Font(size=10, color=GRIS)
    ws["D3"].alignment = Alignment(vertical="center")

    # Línea de marca
    for c in ws["A5:N5"][0]:
        c.fill = PatternFill("solid", fgColor=NARANJA)
    ws.row_dimensions[5].height = 5

    peor = informe.get("peor_global")
    peor_region = informe.get("peor_region")
    top = informe.get("ranking", [])

    kpi(ws, "A7:C9", "AUDITORÍAS ANALIZADAS", str(informe.get("total_auditorias", 0)), "neutral")
    kpi(ws, "D7:F9", "ASESORES CON QA", str(informe.get("total_asesores", 0)), "neutral")
    kpi(ws, "G7:I9", "PEOR PROMEDIO QA", f"{peor['promedio']*100:.1f}%" if peor else "Sin datos", "rojo" if peor else "neutral")
    kpi(ws, "J7:L9", "CASOS CRÍTICOS", str(sum(x.get("criticos", 0) for x in informe.get("ranking_completo", []))), "ambar")
    kpi(ws, "M7:N9", "REGIONES", str(len(informe.get("regiones", []))), "verde")

    # Hallazgo principal
    ws.merge_cells("A11:N11")
    ws["A11"] = "HALLAZGO PRINCIPAL"
    ws["A11"].fill = PatternFill("solid", fgColor=NARANJA)
    ws["A11"].font = Font(color=BLANCO, bold=True, size=11)
    ws["A11"].alignment = Alignment(vertical="center")

    ws.merge_cells("A12:N15")
    if peor:
        muestra = peor.get("estado_muestra", {}).get("texto", "")
        region = peor.get("region", "Sin región")
        principal = (peor.get("peor_item") or {}).get("nombre", "Sin detalle suficiente")
        tendencia = peor.get("tendencia", "Sin tendencia")
        texto = (
            f"{peor['asesor']} presenta la menor calificación del grupo analizado con "
            f"{peor['promedio']*100:.1f}% en {peor['auditorias']} auditorías. "
            f"Región: {region}. Estado de muestra: {muestra}. "
            f"Principal oportunidad detectada: {principal}. "
            f"Tendencia observada: {tendencia}. "
            f"Riesgo QA: {peor.get('riesgo_qa',0):.1f}/100 ({peor.get('nivel_riesgo','')})."
        )
    else:
        texto = "No hay evaluaciones suficientes en el corte seleccionado para identificar un peor desempeño."
    ws["A12"] = texto
    ws["A12"].font = Font(size=11, color=NEGRO)
    ws["A12"].alignment = Alignment(vertical="top", wrap_text=True)
    ws["A12"].fill = PatternFill("solid", fgColor="FFF9F5")
    for row in ws["A12:N15"]:
        for c in row:
            c.border = border

    # Top N
    titulo_top = f"TOP {len(top)} · ASESORES CON MENOR CALIFICACIÓN QA" if top else "ASESORES CON MENOR CALIFICACIÓN QA"
    estilizar_titulo(ws, "A17:H17", titulo_top)
    headers = ["#", "Asesor", "Región", "Auditorías", "Muestra", "Promedio QA", "Críticos", "Principal falla"]
    for col, h in enumerate(headers, 1):
        ws.cell(18, col, h)
    encabezado_tabla(ws, 18, 1, 8)

    row = 19
    for x in top:
        ws.cell(row, 1, x["posicion"])
        ws.cell(row, 2, x["asesor"])
        ws.cell(row, 3, x["region"])
        ws.cell(row, 4, x["auditorias"])
        ws.cell(row, 5, x["estado_muestra"]["texto"])
        ws.cell(row, 6, x["promedio"])
        ws.cell(row, 7, x["criticos"])
        ws.cell(row, 8, (x.get("peor_item") or {}).get("nombre", "Sin detalle"))
        ws.cell(row, 6).number_format = "0.0%"
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=formato_muestra(x["estado_muestra"]["nivel"]))
        row += 1
    if top:
        dar_formato_tabla(ws, 19, row-1, 1, 8)
        ws.conditional_formatting.add(
            f"F19:F{row-1}",
            ColorScaleRule(
                start_type="num", start_value=0, start_color=ROJO,
                mid_type="num", mid_value=0.8, mid_color="F7DC6F",
                end_type="num", end_value=1, end_color="63BE7B"
            )
        )

        # Gráfico Top N
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Top menor calificación QA"
        chart.y_axis.title = "Asesor"
        chart.x_axis.title = "Promedio QA"
        chart.x_axis.scaling.min = 0
        chart.x_axis.scaling.max = 1
        chart.x_axis.numFmt = "0%"
        chart.height = 8.2
        chart.width = 13.5
        data = Reference(ws, min_col=6, min_row=18, max_row=row-1)
        cats = Reference(ws, min_col=2, min_row=19, max_row=row-1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend = None
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = True
        ws.add_chart(chart, "J17")

    # Ranking regional
    regional_start = max(row + 2, 36)
    estilizar_titulo(ws, f"A{regional_start}:F{regional_start}", "DESEMPEÑO POR REGIÓN")
    reg_head = regional_start + 1
    for col, h in enumerate(["Región", "Auditorías", "Asesores", "Promedio QA", "Críticos", "Tasa crítica"], 1):
        ws.cell(reg_head, col, h)
    encabezado_tabla(ws, reg_head, 1, 6)
    rr = reg_head + 1
    for x in informe.get("regiones", []):
        ws.cell(rr, 1, x["region"])
        ws.cell(rr, 2, x["auditorias"])
        ws.cell(rr, 3, x["asesores"])
        ws.cell(rr, 4, x["promedio"])
        ws.cell(rr, 5, x["criticos"])
        ws.cell(rr, 6, x["tasa_criticos"])
        ws.cell(rr, 4).number_format = "0.0%"
        ws.cell(rr, 6).number_format = "0.0%"
        rr += 1
    if rr > reg_head + 1:
        dar_formato_tabla(ws, reg_head+1, rr-1, 1, 6)
        ws.conditional_formatting.add(
            f"D{reg_head+1}:D{rr-1}",
            ColorScaleRule(
                start_type="num", start_value=0, start_color=ROJO,
                mid_type="num", mid_value=0.8, mid_color="F7DC6F",
                end_type="num", end_value=1, end_color="63BE7B"
            )
        )
        chart_reg = BarChart()
        chart_reg.type = "col"
        chart_reg.style = 10
        chart_reg.title = "Calidad promedio por región"
        chart_reg.y_axis.title = "Promedio QA"
        chart_reg.y_axis.scaling.min = 0
        chart_reg.y_axis.scaling.max = 1
        chart_reg.y_axis.numFmt = "0%"
        chart_reg.height = 7.5
        chart_reg.width = 13.5
        data = Reference(ws, min_col=4, min_row=reg_head, max_row=rr-1)
        cats = Reference(ws, min_col=1, min_row=reg_head+1, max_row=rr-1)
        chart_reg.add_data(data, titles_from_data=True)
        chart_reg.set_categories(cats)
        chart_reg.legend = None
        ws.add_chart(chart_reg, f"H{regional_start}")

    # Recomendaciones ejecutivas
    rec_start = max(rr + 2, regional_start + 16)
    estilizar_titulo(ws, f"A{rec_start}:N{rec_start}", "LECTURA EJECUTIVA Y RECOMENDACIONES")
    ws.merge_cells(start_row=rec_start+1, start_column=1, end_row=rec_start+7, end_column=14)
    recomendaciones = []
    if peor:
        recomendaciones.append(
            f"1. Priorizar seguimiento de {peor['asesor']} por registrar el menor promedio QA del corte."
        )
        if peor.get("criticos", 0):
            recomendaciones.append(
                f"2. Revisar {peor['criticos']} caso(s) crítico(s) asociados al asesor antes del siguiente coaching."
            )
        if peor.get("estado_muestra", {}).get("codigo") in {"muy_baja", "baja"}:
            recomendaciones.append(
                "3. La muestra del peor calificado es limitada; aumentar auditorías antes de tomar decisiones definitivas de desempeño."
            )
        if peor.get("peor_item"):
            recomendaciones.append(
                f"4. Enfocar coaching en “{peor['peor_item']['nombre']}”, principal fuente de pérdida detectada."
            )
    if peor_region:
        recomendaciones.append(
            f"5. {peor_region['region']} registra el menor promedio regional ({peor_region['promedio']*100:.1f}%); validar si el patrón es individual o de proceso."
        )
    if not recomendaciones:
        recomendaciones = ["No se identificaron recomendaciones por falta de datos suficientes."]
    ws.cell(rec_start+1, 1, "\n\n".join(recomendaciones))
    ws.cell(rec_start+1, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(rec_start+1, 1).font = Font(size=11, color=NEGRO)
    ws.cell(rec_start+1, 1).fill = PatternFill("solid", fgColor=GRIS_CLARO)
    for r in ws.iter_rows(min_row=rec_start+1, max_row=rec_start+7, min_col=1, max_col=14):
        for c in r:
            c.border = border

    # ---------- Hoja: Ranking completo ----------
    ws_full = wb.create_sheet("Ranking completo")
    ws_full.sheet_view.showGridLines = False
    ws_full.freeze_panes = "A2"
    full_headers = [
        "Posición", "Asesor", "Región", "Bandeja principal", "Auditorías",
        "Estado muestra", "Promedio QA", "Casos críticos", "Tasa críticos",
        "Criterios reincidentes", "Tendencia", "Delta tendencia",
        "Principal categoría de falla", "Principal subítem de falla",
        "Riesgo QA", "Nivel riesgo"
    ]
    ws_full.append(full_headers)
    encabezado_tabla(ws_full, 1, 1, len(full_headers))
    for x in informe.get("ranking_completo", []):
        ws_full.append([
            x["posicion"], x["asesor"], x["region"], x["bandeja_principal"],
            x["auditorias"], x["estado_muestra"]["texto"], x["promedio"],
            x["criticos"], x["tasa_criticos"], x["criterios_reincidentes"],
            x["tendencia"], x["tendencia_delta"],
            (x["peor_categoria"] or {}).get("nombre", ""),
            (x["peor_item"] or {}).get("nombre", ""),
            x["riesgo_qa"], x["nivel_riesgo"],
        ])
    if ws_full.max_row > 1:
        dar_formato_tabla(ws_full, 2, ws_full.max_row, 1, len(full_headers))
        for r in range(2, ws_full.max_row + 1):
            ws_full.cell(r, 7).number_format = "0.0%"
            ws_full.cell(r, 9).number_format = "0.0%"
            if ws_full.cell(r, 12).value is not None:
                ws_full.cell(r, 12).number_format = "+0.0%;-0.0%;0.0%"
        ws_full.conditional_formatting.add(
            f"G2:G{ws_full.max_row}",
            ColorScaleRule(
                start_type="num", start_value=0, start_color=ROJO,
                mid_type="num", mid_value=0.8, mid_color="F7DC6F",
                end_type="num", end_value=1, end_color="63BE7B"
            )
        )
    ws_full.auto_filter.ref = ws_full.dimensions

    # ---------- Hoja: Población seleccionada ----------
    if informe.get("modo") == "seleccion":
        wsp = wb.create_sheet("Población seleccionada")
        wsp.sheet_view.showGridLines = False
        headers_p = ["Asesor solicitado", "Encontrado", "Nombre usado", "Auditorías", "Promedio QA", "Estado muestra", "Posición"]
        wsp.append(headers_p)
        encabezado_tabla(wsp, 1, 1, len(headers_p))
        por_clave = {normalizar_nombre(x["asesor"]): x for x in informe.get("ranking_completo", [])}
        for item in informe.get("resolucion_seleccion", []):
            fila = por_clave.get(normalizar_nombre(item["asesor"]))
            wsp.append([
                item["solicitado"],
                "Sí" if item["encontrado"] else "No",
                item["asesor"] if item["encontrado"] else "",
                fila["auditorias"] if fila else 0,
                fila["promedio"] if fila else None,
                fila["estado_muestra"]["texto"] if fila else "Sin auditorías",
                fila["posicion"] if fila else None,
            ])
        if wsp.max_row > 1:
            dar_formato_tabla(wsp, 2, wsp.max_row, 1, len(headers_p))
            for r in range(2, wsp.max_row + 1):
                if wsp.cell(r, 5).value is not None:
                    wsp.cell(r, 5).number_format = "0.0%"
                if wsp.cell(r, 2).value == "No":
                    for c in range(1, len(headers_p)+1):
                        wsp.cell(r, c).fill = PatternFill("solid", fgColor=ROJO_CLARO)
        wsp.freeze_panes = "A2"
        wsp.auto_filter.ref = wsp.dimensions

    # ---------- Hoja: Regiones ----------
    ws_reg = wb.create_sheet("Regiones")
    ws_reg.sheet_view.showGridLines = False
    reg_headers = ["Región", "Auditorías", "Asesores", "Promedio QA", "Casos críticos", "Tasa críticos", "Peor asesor región", "Promedio peor asesor"]
    ws_reg.append(reg_headers)
    encabezado_tabla(ws_reg, 1, 1, len(reg_headers))
    peor_region_map = {x["region"]: x for x in informe.get("peor_por_region", [])}
    for x in informe.get("regiones", []):
        p = peor_region_map.get(x["region"])
        ws_reg.append([
            x["region"], x["auditorias"], x["asesores"], x["promedio"],
            x["criticos"], x["tasa_criticos"],
            p["asesor"] if p else "",
            p["promedio"] if p else None,
        ])
    if ws_reg.max_row > 1:
        dar_formato_tabla(ws_reg, 2, ws_reg.max_row, 1, len(reg_headers))
        for r in range(2, ws_reg.max_row+1):
            ws_reg.cell(r, 4).number_format = "0.0%"
            ws_reg.cell(r, 6).number_format = "0.0%"
            if ws_reg.cell(r, 8).value is not None:
                ws_reg.cell(r, 8).number_format = "0.0%"
    ws_reg.freeze_panes = "A2"
    ws_reg.auto_filter.ref = ws_reg.dimensions

    # ---------- Hoja: Pérdida de puntos ----------
    ws_loss = wb.create_sheet("Pérdida de puntos")
    ws_loss.sheet_view.showGridLines = False
    loss_headers = ["Subítem", "Categoría", "Veces evaluado", "Casos con falla", "Cumplimiento", "Puntos perdidos"]
    ws_loss.append(loss_headers)
    encabezado_tabla(ws_loss, 1, 1, len(loss_headers))
    for x in informe.get("perdida_global", []):
        ws_loss.append([
            x["nombre"], x["categoria"], x["veces_evaluado"], x["fallas"],
            x["cumplimiento"], x["puntos_perdidos"],
        ])
    if ws_loss.max_row > 1:
        dar_formato_tabla(ws_loss, 2, ws_loss.max_row, 1, len(loss_headers))
        for r in range(2, ws_loss.max_row+1):
            ws_loss.cell(r, 5).number_format = "0.0%"
        chart_loss = BarChart()
        chart_loss.type = "bar"
        chart_loss.style = 10
        chart_loss.title = "Principales pérdidas de puntos"
        chart_loss.height = 8
        chart_loss.width = 14
        data = Reference(ws_loss, min_col=6, min_row=1, max_row=ws_loss.max_row)
        cats = Reference(ws_loss, min_col=1, min_row=2, max_row=ws_loss.max_row)
        chart_loss.add_data(data, titles_from_data=True)
        chart_loss.set_categories(cats)
        chart_loss.legend = None
        ws_loss.add_chart(chart_loss, "H2")
    ws_loss.freeze_panes = "A2"
    ws_loss.auto_filter.ref = ws_loss.dimensions

    # ---------- Hoja: Muestra insuficiente ----------
    if informe.get("muestra_insuficiente"):
        ws_sample = wb.create_sheet("Muestra insuficiente")
        ws_sample.sheet_view.showGridLines = False
        sample_headers = ["Asesor", "Región", "Auditorías", "Promedio QA", "Motivo"]
        ws_sample.append(sample_headers)
        encabezado_tabla(ws_sample, 1, 1, len(sample_headers))
        minimo = informe["filtros"]["minimo_auditorias"]
        for x in informe.get("muestra_insuficiente", []):
            ws_sample.append([
                x["asesor"], x["region"], x["auditorias"], x["promedio"],
                f"Menos de {minimo} auditorías",
            ])
        if ws_sample.max_row > 1:
            dar_formato_tabla(ws_sample, 2, ws_sample.max_row, 1, len(sample_headers))
            for r in range(2, ws_sample.max_row + 1):
                ws_sample.cell(r, 4).number_format = "0.0%"
        ws_sample.freeze_panes = "A2"

    # Ajuste final de anchos en hojas de datos
    for sheet in wb.worksheets[1:]:
        for col in range(1, sheet.max_column + 1):
            ancho = 11
            for cell in sheet[get_column_letter(col)]:
                ancho = max(ancho, min(len(str(cell.value or "")) + 2, 38))
            sheet.column_dimensions[get_column_letter(col)].width = ancho
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Configuración de impresión del resumen
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Dropi QA · Informe interno de calidad"
    ws.oddFooter.right.text = "Página &P de &N"

    wb.save(ruta)
    return ruta

def resumen_terminal(informe):
    peor = informe.get("peor_global")
    peor_region = informe.get("peor_region")
    print("\n" + "=" * 68)
    print("INFORME INTEGRAL DE DESEMPEÑO QA")
    print("=" * 68)
    print(f"Auditorías analizadas: {informe['total_auditorias']}")
    print(f"Asesores encontrados: {informe['total_asesores']}")
    print(f"Mínimo para ranking: {informe['filtros']['minimo_auditorias']} auditorías")
    print()
    if peor:
        print("PEOR CALIFICADO GLOBAL")
        print(f"  Asesor: {peor['asesor']}")
        print(f"  Región: {peor['region']}")
        print(f"  Promedio QA: {peor['promedio']*100:.1f}%")
        print(f"  Auditorías: {peor['auditorias']}")
        print(f"  Casos críticos: {peor['criticos']}")
        print(f"  Criterios reincidentes: {peor['criterios_reincidentes']}")
        print(f"  Tendencia: {peor['tendencia']}")
        if peor["peor_categoria"]:
            print(f"  Principal categoría de falla: {peor['peor_categoria']['nombre']}")
        if peor["peor_item"]:
            print(f"  Principal subítem de falla: {peor['peor_item']['nombre']}")
        print(f"  Riesgo QA: {peor['riesgo_qa']}/100 ({peor['nivel_riesgo']})")
    else:
        print("No hay asesores con muestra suficiente para declarar un peor calificado.")
    print()
    if peor_region:
        print(f"REGIÓN CON MENOR DESEMPEÑO: {peor_region['region']} — {peor_region['promedio']*100:.1f}%")
    print("=" * 68)


def main():
    parser = argparse.ArgumentParser(description="Genera informe de desempeño QA desde el historial local.")
    parser.add_argument("--minimo", type=int, default=10, help="Mínimo de auditorías para ranking oficial.")
    parser.add_argument("--desde", default=None, help="Fecha desde YYYY-MM-DD.")
    parser.add_argument("--hasta", default=None, help="Fecha hasta YYYY-MM-DD.")
    parser.add_argument("--pais", default=None)
    parser.add_argument("--bandeja", default=None)
    parser.add_argument("--excel", default=None, help="Ruta de salida XLSX.")
    args = parser.parse_args()
    informe = construir_informe(
        minimo_auditorias=max(1, args.minimo),
        fecha_desde=args.desde, fecha_hasta=args.hasta,
        pais=args.pais, bandeja=args.bandeja
    )
    resumen_terminal(informe)
    if args.excel:
        ruta = generar_excel(informe, args.excel)
        print(f"\nExcel generado: {ruta.resolve()}")


if __name__ == "__main__":
    main()
